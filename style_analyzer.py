"""
style_analyzer.py
中沢頌の過去文書を読み込み、文体プロファイル (style_profile.md) を生成する。
一度だけ実行すればよい。
"""

import os
import json
from pathlib import Path
from typing import Generator

import anthropic
import pdfplumber
import openpyxl
from docx import Document

from config import (
    DOCS_DIR,
    STYLE_PROFILE_PATH,
    SUPPORTED_EXTENSIONS,
    BATCH_SIZE,
    MAX_CHARS_PER_FILE,
    ANTHROPIC_API_KEY,
    MODEL,
)


# ──────────────────────────────────────────────
# ファイル読み込みユーティリティ
# ──────────────────────────────────────────────

def read_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_pdf(path: str) -> str:
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
    return "\n".join(texts)


def read_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join(str(c) if c is not None else "" for c in row)
            if line.strip():
                rows.append(line)
    return "\n".join(rows)


def read_file(path: str) -> str:
    ext = Path(path).suffix.lower()
    try:
        if ext == ".docx":
            return read_docx(path)
        elif ext == ".pdf":
            return read_pdf(path)
        elif ext == ".xlsx":
            return read_xlsx(path)
        elif ext in (".txt", ".md"):
            with open(path, encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception as e:
        print(f"  [skip] {Path(path).name}: {e}")
    return ""


def collect_files(root: str) -> list[str]:
    """対応拡張子のファイルを再帰的に収集"""
    found = []
    for dirpath, _, filenames in os.walk(root):
        for fn in filenames:
            if Path(fn).suffix.lower() in SUPPORTED_EXTENSIONS:
                found.append(os.path.join(dirpath, fn))
    return sorted(found)


def iter_batches(files: list[str], batch_size: int) -> Generator[list[tuple[str, str]], None, None]:
    """ファイルをバッチに分けてテキスト付きで返す"""
    batch = []
    for path in files:
        text = read_file(path)[:MAX_CHARS_PER_FILE]
        if len(text) > 100:  # 短すぎるものは除外
            batch.append((Path(path).name, text))
            if len(batch) >= batch_size:
                yield batch
                batch = []
    if batch:
        yield batch


# ──────────────────────────────────────────────
# Claude による文体解析
# ──────────────────────────────────────────────

BATCH_ANALYSIS_PROMPT = """\
以下は中沢頌（応用理工学類の大学生）が過去に書いた文書のサンプルです。
これらを注意深く読み、彼の文体・表現・構成の特徴を分析してください。

## 分析する観点
1. **文体・語調** - 常体(だ・である)か敬体(です・ます)か、硬さ・柔らかさ
2. **文の長さ・構造** - 短文/長文の傾向、複文の使い方
3. **専門用語の扱い** - 定義の仕方、略語の使用頻度
4. **論理展開** - 結論先出しか後出しか、根拠の示し方
5. **数式・図表の扱い** - どう引用・説明するか
6. **接続詞・表現の癖** - よく使う接続詞、言い回し
7. **考察の書き方** - どう誤差・例外・限界を述べるか
8. **引用・参考文献スタイル** - 番号式か著者年式か

## サンプル文書

{samples}

---
分析結果をMarkdown形式で詳細に記述してください。"""


AGGREGATE_PROMPT = """\
以下は複数のバッチ分析から得られた「中沢頌の文体特徴」のまとめです。
これらを統合し、一貫した包括的な文体プロファイルを作成してください。

最終プロファイルは以下のセクションを含むMarkdown形式で出力してください：
1. 総合的な文体の印象（3〜5文）
2. 語調・文体（常体/敬体、硬さ、etc.）
3. 文の構造・長さの傾向
4. 専門用語・記号の扱い
5. 論理展開・構成パターン
6. 数式・図・表の引用スタイル
7. よく使う接続詞・表現集（箇条書き、例文付き）
8. 考察・議論の書き方
9. 参考文献スタイル
10. 再現時の注意点・NG表現

## バッチ分析結果

{batch_results}
"""


def analyze_batch(client: anthropic.Anthropic, batch: list[tuple[str, str]]) -> str:
    samples = ""
    for name, text in batch:
        samples += f"\n### {name}\n```\n{text}\n```\n"

    response = client.messages.create(
        model=MODEL,
        max_tokens=2048,
        messages=[{"role": "user", "content": BATCH_ANALYSIS_PROMPT.format(samples=samples)}],
    )
    return response.content[0].text


def aggregate_analyses(client: anthropic.Anthropic, batch_results: list[str]) -> str:
    combined = "\n\n---\n\n".join(
        f"## バッチ {i+1}\n{r}" for i, r in enumerate(batch_results)
    )
    response = client.messages.create(
        model=MODEL,
        max_tokens=4096,
        messages=[{"role": "user", "content": AGGREGATE_PROMPT.format(batch_results=combined)}],
    )
    return response.content[0].text


# ──────────────────────────────────────────────
# メイン実行
# ──────────────────────────────────────────────

def run(docs_dir: str | None = None):
    docs_dir = docs_dir or DOCS_DIR

    if not os.path.isdir(docs_dir):
        print(f"[ERROR] ドキュメントフォルダが見つかりません: {docs_dir}")
        print("  DOCS_DIR 環境変数で正しいパスを指定してください。")
        return

    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    print(f"[1/3] ファイルを収集中: {docs_dir}")
    files = collect_files(docs_dir)
    print(f"  → {len(files)} ファイル見つかりました")

    if not files:
        print("[ERROR] 対応ファイルが見つかりません。")
        return

    print(f"[2/3] バッチ解析中（バッチサイズ={BATCH_SIZE}）...")
    batch_results = []
    for i, batch in enumerate(iter_batches(files, BATCH_SIZE)):
        print(f"  バッチ {i+1}: {[b[0] for b in batch]}")
        result = analyze_batch(client, batch)
        batch_results.append(result)
        print(f"  → 完了")

    if not batch_results:
        print("[ERROR] 解析できる内容がありませんでした。")
        return

    print("[3/3] 文体プロファイルを統合中...")
    if len(batch_results) == 1:
        profile = batch_results[0]
    else:
        profile = aggregate_analyses(client, batch_results)

    header = "# 中沢頌 文体プロファイル\n\n"
    header += "> このファイルは style_analyzer.py によって自動生成されました。\n\n"
    full_profile = header + profile

    with open(STYLE_PROFILE_PATH, "w", encoding="utf-8") as f:
        f.write(full_profile)

    print(f"\n✅ 文体プロファイルを保存しました: {STYLE_PROFILE_PATH}")


if __name__ == "__main__":
    import sys
    docs_dir = sys.argv[1] if len(sys.argv) > 1 else None
    run(docs_dir)
