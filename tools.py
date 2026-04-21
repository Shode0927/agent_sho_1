"""
tools.py
レポートエージェントが使用するツール群。
- run_python      : 数値計算・Python実行
- web_search      : 参考文献・情報検索
- generate_graph  : matplotlib グラフ生成
- format_latex    : LaTeX数式フォーマット
- write_document  : Word / PDF 出力
"""

import io
import os
import sys
import json
import textwrap
import traceback
import base64
from contextlib import redirect_stdout, redirect_stderr
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np
import sympy as sp
from duckduckgo_search import DDGS
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage,
    Table, TableStyle, PageBreak,
)
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from config import OUTPUT_DIR

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ──────────────────────────────────────────────
# Tool 1: Python 実行（数値計算）
# ──────────────────────────────────────────────

TOOL_RUN_PYTHON = {
    "name": "run_python",
    "description": (
        "Pythonコードを実行して数値計算・データ処理を行います。"
        "numpy, scipy, sympy, matplotlib が使えます。"
        "計算結果、変数の値、生成したグラフのファイルパスを返します。"
        "グラフを生成する場合は plt.savefig(filepath) を使い、filepathを出力に含めてください。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "code": {
                "type": "string",
                "description": "実行するPythonコード"
            },
            "description": {
                "type": "string",
                "description": "このコードが何を計算・実行するかの説明（日本語）"
            }
        },
        "required": ["code", "description"]
    }
}


def run_python(code: str, description: str) -> dict:
    stdout_buf = io.StringIO()
    stderr_buf = io.StringIO()

    # グラフ保存用のディレクトリを渡す
    exec_globals = {
        "np": np,
        "sp": sp,
        "plt": plt,
        "OUTPUT_DIR": OUTPUT_DIR,
        "__builtins__": __builtins__,
    }

    try:
        import scipy
        exec_globals["scipy"] = scipy
    except ImportError:
        pass

    try:
        with redirect_stdout(stdout_buf), redirect_stderr(stderr_buf):
            exec(code, exec_globals)  # noqa: S102
        stdout = stdout_buf.getvalue()
        stderr = stderr_buf.getvalue()

        result = {"status": "success", "stdout": stdout}
        if stderr:
            result["stderr"] = stderr

        # 生成されたグラフを検出
        graphs = [
            str(p) for p in Path(OUTPUT_DIR).glob("*.png")
            if p.stat().st_mtime > (datetime.now().timestamp() - 10)
        ]
        if graphs:
            result["generated_graphs"] = graphs

        return result

    except Exception:
        return {
            "status": "error",
            "error": traceback.format_exc(),
            "stdout": stdout_buf.getvalue(),
        }


# ──────────────────────────────────────────────
# Tool 2: Web 検索
# ──────────────────────────────────────────────

TOOL_WEB_SEARCH = {
    "name": "web_search",
    "description": (
        "DuckDuckGoで学術情報・参考文献を検索します。"
        "論文タイトル、著者名、概念の定義、公式の出典などを調べるために使います。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "query": {
                "type": "string",
                "description": "検索クエリ（英語または日本語）"
            },
            "max_results": {
                "type": "integer",
                "description": "取得する結果の最大数（デフォルト5）",
                "default": 5
            }
        },
        "required": ["query"]
    }
}


def web_search(query: str, max_results: int = 5) -> dict:
    try:
        results = []
        with DDGS() as ddgs:
            for r in ddgs.text(query, max_results=max_results):
                results.append({
                    "title": r.get("title", ""),
                    "url": r.get("href", ""),
                    "snippet": r.get("body", "")
                })
        return {"status": "success", "results": results}
    except Exception as e:
        return {"status": "error", "error": str(e)}


# ──────────────────────────────────────────────
# Tool 3: グラフ生成
# ──────────────────────────────────────────────

TOOL_GENERATE_GRAPH = {
    "name": "generate_graph",
    "description": (
        "matplotlib を使ってグラフを生成し、PNGファイルとして保存します。"
        "折れ線グラフ、散布図、棒グラフ、ヒストグラム、等高線図などが作れます。"
        "生成したファイルのパスを返します。write_document で挿入できます。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "graph_type": {
                "type": "string",
                "enum": ["line", "scatter", "bar", "histogram", "errorbar", "contour", "custom"],
                "description": "グラフの種類"
            },
            "data": {
                "type": "object",
                "description": (
                    "グラフデータ。例: "
                    "{\"x\": [1,2,3], \"y\": [4,5,6], \"xlabel\": \"時間 [s]\", "
                    "\"ylabel\": \"変位 [m]\", \"title\": \"変位-時間グラフ\", "
                    "\"label\": \"実験値\"}"
                )
            },
            "filename": {
                "type": "string",
                "description": "保存ファイル名（拡張子なし）"
            }
        },
        "required": ["graph_type", "data", "filename"]
    }
}


def generate_graph(graph_type: str, data: dict, filename: str) -> dict:
    try:
        fig, ax = plt.subplots(figsize=(8, 5), dpi=150)

        # 日本語フォント設定（環境依存で fallback）
        try:
            plt.rcParams["font.family"] = "IPAexGothic"
        except Exception:
            pass

        x = data.get("x", [])
        y = data.get("y", [])
        label = data.get("label", "")
        color = data.get("color", "royalblue")

        if graph_type == "line":
            ax.plot(x, y, marker="o", label=label, color=color)
        elif graph_type == "scatter":
            ax.scatter(x, y, label=label, color=color)
        elif graph_type == "bar":
            ax.bar(x, y, label=label, color=color)
        elif graph_type == "histogram":
            ax.hist(y, bins=data.get("bins", 20), label=label, color=color)
        elif graph_type == "errorbar":
            yerr = data.get("yerr", None)
            ax.errorbar(x, y, yerr=yerr, fmt="o-", label=label, color=color,
                        capsize=4, elinewidth=1.5)
        elif graph_type == "contour":
            X = np.array(x)
            Y = np.array(y)
            Z = np.array(data.get("z", []))
            cs = ax.contourf(X, Y, Z, levels=data.get("levels", 20), cmap="viridis")
            fig.colorbar(cs, ax=ax, label=data.get("zlabel", ""))
        elif graph_type == "custom":
            # カスタム: run_python で matplotlib を直接使う方を推奨
            pass

        ax.set_xlabel(data.get("xlabel", ""), fontsize=12)
        ax.set_ylabel(data.get("ylabel", ""), fontsize=12)
        ax.set_title(data.get("title", ""), fontsize=13)
        if label:
            ax.legend(fontsize=10)
        ax.grid(True, linestyle="--", alpha=0.6)

        # 追加のデータセット（y2など）
        if "y2" in data:
            ax.plot(x, data["y2"], marker="s", label=data.get("label2", ""), color="tomato")
            ax.legend(fontsize=10)

        plt.tight_layout()
        filepath = os.path.join(OUTPUT_DIR, f"{filename}.png")
        plt.savefig(filepath)
        plt.close(fig)

        return {"status": "success", "filepath": filepath}

    except Exception as e:
        return {"status": "error", "error": traceback.format_exc()}


# ──────────────────────────────────────────────
# Tool 4: LaTeX 数式フォーマット
# ──────────────────────────────────────────────

TOOL_FORMAT_LATEX = {
    "name": "format_latex",
    "description": (
        "数式をLaTeX形式に変換・検証します。"
        "sympy を使って数式を正規化し、LaTeXコードを返します。"
        "Word/PDFに埋め込む際の表記を確定するために使います。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "expression": {
                "type": "string",
                "description": "LaTeX数式または sympy 式（例: 'x**2 + 2*x + 1' または '\\\\frac{d^2x}{dt^2}'）"
            },
            "mode": {
                "type": "string",
                "enum": ["sympy_to_latex", "validate_latex", "simplify"],
                "description": "sympy_to_latex: sympyをLaTeXに変換, validate_latex: LaTeX構文チェック, simplify: 式を簡略化"
            }
        },
        "required": ["expression", "mode"]
    }
}


def format_latex(expression: str, mode: str) -> dict:
    try:
        if mode == "sympy_to_latex":
            expr = sp.sympify(expression)
            latex_str = sp.latex(expr)
            return {"status": "success", "latex": latex_str, "sympy": str(expr)}

        elif mode == "simplify":
            expr = sp.sympify(expression)
            simplified = sp.simplify(expr)
            latex_str = sp.latex(simplified)
            return {
                "status": "success",
                "original": str(expr),
                "simplified": str(simplified),
                "latex": latex_str
            }

        elif mode == "validate_latex":
            # 基本的な構文チェック（括弧の対応など）
            depth = 0
            for ch in expression:
                if ch == "{":
                    depth += 1
                elif ch == "}":
                    depth -= 1
                if depth < 0:
                    return {"status": "error", "error": "不正な括弧の対応"}
            if depth != 0:
                return {"status": "error", "error": f"閉じていない括弧が {depth} 個あります"}
            return {"status": "success", "latex": expression, "message": "構文チェックOK"}

    except Exception as e:
        return {"status": "error", "error": str(e)}


# ──────────────────────────────────────────────
# Tool 5: ドキュメント出力 (Word / PDF)
# ──────────────────────────────────────────────

TOOL_WRITE_DOCUMENT = {
    "name": "write_document",
    "description": (
        "レポートをWordまたはPDFファイルとして出力します。"
        "sections に各セクションを渡すと、見出し・本文・数式・画像を整形して出力します。"
        "生成したファイルのパスを返します。"
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "filename": {
                "type": "string",
                "description": "出力ファイル名（拡張子なし）"
            },
            "format": {
                "type": "string",
                "enum": ["docx", "pdf"],
                "description": "出力形式"
            },
            "title": {
                "type": "string",
                "description": "レポートタイトル"
            },
            "author": {
                "type": "string",
                "description": "著者名（デフォルト: 中沢頌）"
            },
            "date": {
                "type": "string",
                "description": "日付（省略時は今日）"
            },
            "sections": {
                "type": "array",
                "description": "セクションのリスト",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string", "description": "見出し（省略可）"},
                        "level": {"type": "integer", "description": "見出しレベル 1〜3"},
                        "body": {"type": "string", "description": "本文テキスト"},
                        "latex_equations": {
                            "type": "array",
                            "description": "挿入するLaTeX数式のリスト（文中に [EQ:数式] で参照）",
                            "items": {"type": "string"}
                        },
                        "image_paths": {
                            "type": "array",
                            "description": "挿入する画像のファイルパスリスト",
                            "items": {"type": "string"}
                        },
                        "image_captions": {
                            "type": "array",
                            "description": "画像のキャプションリスト",
                            "items": {"type": "string"}
                        }
                    }
                }
            }
        },
        "required": ["filename", "format", "title", "sections"]
    }
}


def _write_docx(filepath: str, title: str, author: str, date: str, sections: list) -> str:
    doc = Document()

    # スタイル設定
    style = doc.styles["Normal"]
    style.font.name = "游明朝"
    style.font.size = Pt(10.5)

    # タイトル
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    # 著者・日付
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"{author}　{date}").font.size = Pt(10)
    doc.add_paragraph()

    eq_counter = [0]

    for sec in sections:
        heading = sec.get("heading", "")
        level = sec.get("level", 1)
        body = sec.get("body", "")
        equations = sec.get("latex_equations", [])
        image_paths = sec.get("image_paths", [])
        captions = sec.get("image_captions", [])

        if heading:
            doc.add_heading(heading, level=level)

        if body:
            # 数式プレースホルダーを処理
            parts = body.split("[EQ:")
            first = True
            para = None
            for part in parts:
                if first:
                    if part.strip():
                        para = doc.add_paragraph(part)
                    first = False
                else:
                    eq_idx_str, rest = part.split("]", 1) if "]" in part else (part, "")
                    try:
                        eq_idx = int(eq_idx_str.strip()) - 1
                        eq_latex = equations[eq_idx] if eq_idx < len(equations) else eq_idx_str
                    except (ValueError, IndexError):
                        eq_latex = eq_idx_str
                    eq_counter[0] += 1
                    eq_para = doc.add_paragraph()
                    eq_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    eq_run = eq_para.add_run(f"  {eq_latex}  … ({eq_counter[0]})")
                    eq_run.font.name = "Cambria Math"
                    if rest.strip():
                        doc.add_paragraph(rest.strip())

            if first and body.strip() and "[EQ:" not in body:
                doc.add_paragraph(body)

        for i, img_path in enumerate(image_paths):
            if os.path.isfile(img_path):
                try:
                    doc.add_picture(img_path, width=Inches(5))
                    caption_text = captions[i] if i < len(captions) else f"図{eq_counter[0]+1}"
                    cap_para = doc.add_paragraph(caption_text)
                    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = cap_para.runs[0]
                    cap_run.font.size = Pt(9)
                except Exception as e:
                    doc.add_paragraph(f"[画像挿入エラー: {e}]")

    doc.save(filepath)
    return filepath


def _write_pdf(filepath: str, title: str, author: str, date: str, sections: list) -> str:
    # 日本語フォント登録（利用可能なものを試みる）
    jp_font = None
    font_candidates = [
        "/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf",
        "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    ]
    for fc in font_candidates:
        if os.path.isfile(fc):
            try:
                pdfmetrics.registerFont(TTFont("JPFont", fc))
                jp_font = "JPFont"
                break
            except Exception:
                continue

    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        leftMargin=25 * mm,
        rightMargin=25 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm,
    )

    styles = getSampleStyleSheet()
    base_font = jp_font or "Helvetica"

    title_style = ParagraphStyle(
        "ReportTitle",
        fontName=base_font,
        fontSize=16,
        leading=22,
        alignment=1,
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        fontName=base_font,
        fontSize=10,
        alignment=1,
        spaceAfter=12,
    )
    h1_style = ParagraphStyle(
        "H1", fontName=base_font, fontSize=13, leading=18,
        spaceBefore=12, spaceAfter=4, textColor=colors.black,
    )
    h2_style = ParagraphStyle(
        "H2", fontName=base_font, fontSize=11, leading=16,
        spaceBefore=8, spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "Body", fontName=base_font, fontSize=10, leading=16,
        spaceAfter=6,
    )
    eq_style = ParagraphStyle(
        "Equation", fontName="Helvetica", fontSize=10, leading=16,
        alignment=1, spaceAfter=6,
    )
    caption_style = ParagraphStyle(
        "Caption", fontName=base_font, fontSize=9, leading=14,
        alignment=1, spaceAfter=8,
    )

    story = []
    story.append(Paragraph(title, title_style))
    story.append(Paragraph(f"{author}　{date}", subtitle_style))
    story.append(Spacer(1, 6 * mm))

    eq_counter = [0]

    for sec in sections:
        heading = sec.get("heading", "")
        level = sec.get("level", 1)
        body = sec.get("body", "")
        equations = sec.get("latex_equations", [])
        image_paths = sec.get("image_paths", [])
        captions = sec.get("image_captions", [])

        if heading:
            h_style = h1_style if level == 1 else h2_style
            story.append(Paragraph(heading, h_style))

        if body:
            parts = body.split("[EQ:")
            first = True
            for part in parts:
                if first:
                    if part.strip():
                        story.append(Paragraph(part.replace("\n", "<br/>"), body_style))
                    first = False
                else:
                    eq_idx_str, rest = part.split("]", 1) if "]" in part else (part, "")
                    try:
                        eq_idx = int(eq_idx_str.strip()) - 1
                        eq_latex = equations[eq_idx] if eq_idx < len(equations) else eq_idx_str
                    except (ValueError, IndexError):
                        eq_latex = eq_idx_str
                    eq_counter[0] += 1
                    story.append(Paragraph(f"{eq_latex}  …({eq_counter[0]})", eq_style))
                    if rest.strip():
                        story.append(Paragraph(rest.strip().replace("\n", "<br/>"), body_style))

            if first and body.strip() and "[EQ:" not in body:
                story.append(Paragraph(body.replace("\n", "<br/>"), body_style))

        for i, img_path in enumerate(image_paths):
            if os.path.isfile(img_path):
                try:
                    img = RLImage(img_path, width=130 * mm, height=90 * mm, kind="proportional")
                    story.append(img)
                    caption_text = captions[i] if i < len(captions) else f"図{eq_counter[0]+1}"
                    story.append(Paragraph(caption_text, caption_style))
                except Exception as e:
                    story.append(Paragraph(f"[画像挿入エラー: {e}]", body_style))

    doc.build(story)
    return filepath


def write_document(
    filename: str,
    format: str,
    title: str,
    sections: list,
    author: str = "中沢頌",
    date: str = "",
) -> dict:
    if not date:
        date = datetime.now().strftime("%Y年%m月%d日")

    ext = "docx" if format == "docx" else "pdf"
    filepath = os.path.join(OUTPUT_DIR, f"{filename}.{ext}")

    try:
        if format == "docx":
            _write_docx(filepath, title, author, date, sections)
        else:
            _write_pdf(filepath, title, author, date, sections)

        return {"status": "success", "filepath": filepath}
    except Exception as e:
        return {"status": "error", "error": traceback.format_exc()}


# ──────────────────────────────────────────────
# ツールディスパッチャ
# ──────────────────────────────────────────────

ALL_TOOLS = [
    TOOL_RUN_PYTHON,
    TOOL_WEB_SEARCH,
    TOOL_GENERATE_GRAPH,
    TOOL_FORMAT_LATEX,
    TOOL_WRITE_DOCUMENT,
]


def dispatch_tool(tool_name: str, tool_input: dict) -> str:
    """ツール名と入力を受け取り、結果をJSON文字列で返す"""
    if tool_name == "run_python":
        result = run_python(**tool_input)
    elif tool_name == "web_search":
        result = web_search(**tool_input)
    elif tool_name == "generate_graph":
        result = generate_graph(**tool_input)
    elif tool_name == "format_latex":
        result = format_latex(**tool_input)
    elif tool_name == "write_document":
        result = write_document(**tool_input)
    else:
        result = {"status": "error", "error": f"Unknown tool: {tool_name}"}

    return json.dumps(result, ensure_ascii=False)
